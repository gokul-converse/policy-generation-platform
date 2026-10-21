
import csv
import sqlite3
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


BASE_DIR = Path(__file__).resolve().parents[2]
DB_PATH = BASE_DIR / "data" / "enterprise_policy.db"


def extract_text(file_path: Path) -> list[dict]:
    """Extract text while preserving basic page or section information."""

    suffix = file_path.suffix.lower()
    sections = []

    if suffix == ".txt":
        text = file_path.read_text(encoding="utf-8")
        if text.strip():
            sections.append({
                "section": file_path.name,
                "page": None,
                "text": text,
            })

    elif suffix == ".pdf":
        reader = PdfReader(str(file_path))

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            if text.strip():
                sections.append({
                    "section": f"Page {page_number}",
                    "page": page_number,
                    "text": text,
                })

    elif suffix == ".docx":
        doc = Document(str(file_path))
        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )

        if text.strip():
            sections.append({
                "section": "Document body",
                "page": None,
                "text": text,
            })

    elif suffix == ".xlsx":
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            for sheet in workbook.worksheets:
                rows = []

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        rows.append(" | ".join(values))

                if rows:
                    sections.append({
                        "section": sheet.title,
                        "page": None,
                        "text": "\n".join(rows),
                    })
        finally:
            workbook.close()

    elif suffix == ".csv":
        with file_path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:
            rows = [
                " | ".join(row)
                for row in csv.reader(file)
                if any(value.strip() for value in row)
            ]

        if rows:
            sections.append({
                "section": "CSV data",
                "page": None,
                "text": "\n".join(rows),
            })

    else:
        raise ValueError(f"Unsupported file format: {suffix}")

    return sections


def split_text(
    text: str,
    chunk_size: int = 1500,
) -> list[str]:
    """Split extracted text into fixed-size word chunks."""

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    words = text.split()
    chunks = []

    for start in range(0, len(words), chunk_size):
        chunk = " ".join(words[start:start + chunk_size])

        if chunk:
            chunks.append(chunk)

    return chunks


def ingest_document(
    file_path: str,
    domain: str,
    document_type: str | None = None,
    approval_status: str = "sample_not_approved",
    version: str = "1",
    effective_date: str | None = None,
) -> dict:
    """
    Extract a document and save its metadata and chunks in SQLite.

    document_type and effective_date are retained as optional parameters
    for compatibility with existing callers. The simplified database
    does not contain columns for these fields.
    """

    path = Path(file_path).resolve()

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    if not DB_PATH.is_file():
        raise FileNotFoundError(f"Database not found: {DB_PATH}")

    if domain not in {"hr", "finance", "it_security", "general"}:
        raise ValueError(f"Unsupported document domain: {domain}")

    if approval_status not in {
        "approved",
        "pending_review",
        "rejected",
        "sample_not_approved",
        "unknown",
    }:
        raise ValueError(f"Unsupported approval status: {approval_status}")

    # Extract before inserting anything into the database.
    sections = extract_text(path)

    chunks_to_save = []
    chunk_index = 0

    for section in sections:
        for chunk in split_text(section["text"]):
            chunks_to_save.append((
                chunk_index,
                section["section"],
                section["page"],
                chunk,
            ))
            chunk_index += 1

    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("PRAGMA foreign_keys = ON")

        # Avoid inserting the same file path and version twice.
        existing = conn.execute(
            """
            SELECT document_id, ingestion_status
            FROM documents
            WHERE file_path = ? AND version = ?
            """,
            (str(path), version),
        ).fetchone()

        if existing:
            return {
                "document_id": existing[0],
                "document_name": path.name,
                "domain": domain,
                "status": "already_exists",
                "message": (
                    "This file path and version are already registered. "
                    "No duplicate record was created."
                ),
            }

        # Insert document metadata and chunks in one transaction.
        cursor = conn.execute(
            """
            INSERT INTO documents (
                document_name,
                domain,
                version,
                approval_status,
                file_path,
                ingestion_status
            )
            VALUES (?, ?, ?, ?, ?, 'processing')
            """,
            (
                path.name,
                domain,
                version,
                approval_status,
                str(path),
            ),
        )

        document_id = cursor.lastrowid

        conn.executemany(
            """
            INSERT INTO document_chunks (
                document_id,
                chunk_index,
                section_title,
                page_number,
                chunk_text
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            [
                (
                    document_id,
                    index,
                    section,
                    page,
                    text,
                )
                for index, section, page, text in chunks_to_save
            ],
        )

        conn.execute(
            """
            UPDATE documents
            SET ingestion_status = 'completed'
            WHERE document_id = ?
            """,
            (document_id,),
        )

    return {
        "document_id": document_id,
        "document_name": path.name,
        "domain": domain,
        "sections_extracted": len(sections),
        "chunks_saved": len(chunks_to_save),
        "status": "completed",
    }
